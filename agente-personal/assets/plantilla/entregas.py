"""Entregas: Claude deja un .md y un .json en entregas/; este código (de confianza) los convierte y los entrega.

Claude nunca genera binarios ni sube archivos por su cuenta. Destinos (config.json > entregas.destino):
  nextcloud  sube por WebDAV (credenciales en .env: NEXTCLOUD_URL, NEXTCLOUD_USUARIO, NEXTCLOUD_CLAVE)
  carpeta    copia a una carpeta sincronizada (Google Drive, OneDrive, Dropbox...)
  local      deja el archivo en entregas/hechas
"""
import datetime
import json
import os
import re
import shutil
import subprocess
from pathlib import Path
from urllib.parse import quote

import requests

AQUI = Path(__file__).resolve().parent
# Dentro de la carpeta de trabajo: es el único lugar donde Claude puede escribir.
CARPETA = Path(json.loads((AQUI / "config.json").read_text(encoding="utf-8"))["carpeta_trabajo"]).expanduser() / "_entregas"
HECHAS = CARPETA / "hechas"
FORMATOS = {"docx", "pdf", "md"}
SIN_VENTANA = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0


def _config():
    return json.loads((AQUI / "config.json").read_text(encoding="utf-8")).get("entregas", {"destino": "local"})


def instrucciones():
    """Texto para Claude según el destino configurado."""
    destino = _config().get("destino", "local")
    donde = {
        "nextcloud": '"carpeta": "Carpeta/Subcarpeta en Nextcloud"',
        "carpeta": '"carpeta": "Subcarpeta dentro de la carpeta de entregas"',
        "local": '"carpeta": ""',
    }[destino]
    return (
        "ENTREGAS (Word, PDF o documento para guardar; solo en modo trabajo): tú no generas binarios ni subes archivos. "
        f"Escribe el contenido completo en Markdown en {CARPETA}{os.sep}<nombre-corto>.md y, al lado, "
        f"{CARPETA}{os.sep}<nombre-corto>.json con este contenido: "
        '{"archivo": "<nombre-corto>.md", "titulo": "Título legible", "formato": "docx" o "pdf" o "md", '
        f"{donde}}}. Apenas termines, el bot lo convierte, lo entrega y avisa; en tu respuesta di en 3 a 5 líneas "
        "qué contiene."
    )


# ------------------------------------------------------------------ Markdown -> Word -> PDF
def _texto_con_formato(parrafo, texto):
    texto = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1 (\2)", texto)
    texto = texto.replace("`", "")
    for trozo in re.split(r"(\*\*[^*]+\*\*)", texto):
        if trozo.startswith("**") and trozo.endswith("**") and len(trozo) > 4:
            parrafo.add_run(trozo[2:-2]).bold = True
        elif trozo:
            parrafo.add_run(trozo)


def md_a_docx(md, destino, titulo):
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(titulo, level=0)
    fecha = doc.add_paragraph(datetime.date.today().strftime("%d/%m/%Y"))
    fecha.runs[0].italic = True

    for linea in Path(md).read_text(encoding="utf-8").splitlines():
        limpia = linea.strip()
        if not limpia or limpia.startswith("|---") or limpia.startswith("| ---") or limpia == "---":
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", limpia)
        if m:
            nivel = len(m.group(1))
            if nivel == 1 and m.group(2).strip() == titulo:
                continue
            doc.add_heading(m.group(2).strip(), level=min(nivel, 3))
        elif re.match(r"^[-*•]\s+", limpia):
            _texto_con_formato(doc.add_paragraph(style="List Bullet"), re.sub(r"^[-*•]\s+", "", limpia))
        elif re.match(r"^\d+[.)]\s+", limpia):
            _texto_con_formato(doc.add_paragraph(style="List Number"), re.sub(r"^\d+[.)]\s+", "", limpia))
        elif limpia.startswith(">"):
            _texto_con_formato(doc.add_paragraph(style="Intense Quote"), limpia.lstrip("> "))
        elif limpia.startswith("|"):
            celdas = [c.strip() for c in limpia.strip("|").split("|")]
            _texto_con_formato(doc.add_paragraph(), " · ".join(celdas))
        else:
            _texto_con_formato(doc.add_paragraph(), limpia)
    doc.save(destino)


_PS_WORD = (
    "$ErrorActionPreference='Stop'; $w=New-Object -ComObject Word.Application; $w.Visible=$false; "
    "try {{ $d=$w.Documents.Open('{docx}', $false, $true); $d.SaveAs([ref]'{pdf}', [ref]17); $d.Close([ref]$false) }} "
    "finally {{ $w.Quit() }}"
)


def docx_a_pdf(docx, pdf):
    """Word en Windows; si no, LibreOffice en cualquier sistema."""
    docx, pdf = Path(docx).resolve(), Path(pdf).resolve()
    errores = []
    if os.name == "nt":
        script = _PS_WORD.format(docx=str(docx).replace("'", "''"), pdf=str(pdf).replace("'", "''"))
        p = subprocess.run(["powershell", "-NoProfile", "-Command", script], capture_output=True, text=True,
                           timeout=180, creationflags=SIN_VENTANA)
        if pdf.exists():
            return
        errores.append("Word: " + (p.stderr or p.stdout).strip()[-200:])
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    for posible in (r"C:\Program Files\LibreOffice\program\soffice.exe", "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if not soffice and Path(posible).exists():
            soffice = posible
    if soffice:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(pdf.parent), str(docx)],
                       capture_output=True, timeout=180, creationflags=SIN_VENTANA)
        generado = pdf.parent / (docx.stem + ".pdf")
        if generado.exists():
            if generado != pdf:
                generado.replace(pdf)
            return
        errores.append("LibreOffice no generó el PDF")
    raise RuntimeError("no pude crear el PDF (instala Microsoft Word o LibreOffice). " + " | ".join(errores))


# ------------------------------------------------------------------ destinos
def _env():
    datos = {}
    ruta = AQUI / ".env"
    if ruta.exists():
        for linea in ruta.read_text(encoding="utf-8-sig").splitlines():
            if "=" in linea and not linea.strip().startswith("#"):
                k, v = linea.split("=", 1)
                datos[k.strip()] = v.strip().strip('"').strip("'")
    return datos


def subir_nextcloud(archivo, carpeta):
    env = _env()
    host = env.get("NEXTCLOUD_URL", "").rstrip("/")
    usuario, clave = env.get("NEXTCLOUD_USUARIO"), env.get("NEXTCLOUD_CLAVE")
    if not (host and usuario and clave):
        raise RuntimeError("faltan NEXTCLOUD_URL, NEXTCLOUD_USUARIO o NEXTCLOUD_CLAVE en .env (corre configurar.py)")
    if not host.startswith("http"):
        host = "https://" + host
    base = f"{host}/remote.php/dav/files/{quote(usuario)}"
    sesion = requests.Session()
    sesion.auth = (usuario, clave)
    ruta = ""
    for parte in [p for p in carpeta.split("/") if p]:
        ruta += "/" + quote(parte)
        r = sesion.request("MKCOL", base + ruta, timeout=60)
        if r.status_code not in (201, 405):
            raise RuntimeError(f"Nextcloud no dejó crear la carpeta ({r.status_code})")
    with open(archivo, "rb") as f:
        r = sesion.put(f"{base}{ruta}/{quote(Path(archivo).name)}", data=f, timeout=300)
    if r.status_code not in (200, 201, 204):
        raise RuntimeError(f"Nextcloud rechazó el archivo ({r.status_code})")
    return f"📎 Subí «{Path(archivo).name}» a Nextcloud › {carpeta}\n{host}/index.php/apps/files/?dir={quote('/' + carpeta)}"


def copiar_a_carpeta(archivo, carpeta):
    raiz = Path(_config().get("carpeta", "")).expanduser()
    if not str(raiz) or not raiz.exists():
        raise RuntimeError(f"la carpeta de entregas no existe: {raiz}")
    destino = (raiz / carpeta).resolve()
    if raiz.resolve() not in (destino, *destino.parents):
        raise RuntimeError("la subcarpeta se sale de la carpeta de entregas")
    destino.mkdir(parents=True, exist_ok=True)
    final = destino / Path(archivo).name
    shutil.copy2(archivo, final)
    return f"📎 Guardé «{final.name}» en {final.parent}"


# ------------------------------------------------------------------ orquestación
def _nombre_seguro(texto):
    limpio = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", texto).strip().rstrip(".")
    return limpio[:120] or "Documento"


def _validar(pedido, json_path, destino):
    md = (CARPETA / str(pedido.get("archivo", ""))).resolve()
    if md.parent != CARPETA.resolve() or md.suffix.lower() != ".md" or not md.exists():
        raise ValueError(f"el Markdown indicado no existe en entregas/ ({pedido.get('archivo')})")
    formato = str(pedido.get("formato", "docx")).lower()
    if formato not in FORMATOS:
        raise ValueError(f"formato no soportado: {formato}")
    carpeta = str(pedido.get("carpeta") or pedido.get("carpeta_nc") or "").replace("\\", "/").strip("/")
    if ".." in carpeta.split("/"):
        raise ValueError("carpeta no válida")
    if destino == "nextcloud" and not carpeta:
        raise ValueError("falta la carpeta de Nextcloud")
    titulo = _nombre_seguro(str(pedido.get("titulo") or json_path.stem))
    return md, formato, carpeta, titulo


def procesar_pendientes(log=None):
    """Convierte y entrega cada pedido pendiente. Devuelve líneas para contarle a la persona."""
    CARPETA.mkdir(exist_ok=True)
    destino = _config().get("destino", "local")
    avisos = []
    for pedido_json in sorted(CARPETA.glob("*.json")):
        guardar_en = HECHAS / datetime.datetime.now().strftime("%Y-%m-%d_%H%M%S")
        try:
            pedido = json.loads(pedido_json.read_text(encoding="utf-8"))
            md, formato, carpeta, titulo = _validar(pedido, pedido_json, destino)
            guardar_en.mkdir(parents=True, exist_ok=True)
            if formato == "md":
                final = guardar_en / f"{titulo}.md"
                shutil.copy2(md, final)
            else:
                final = guardar_en / f"{titulo}.docx"
                md_a_docx(md, final, titulo)
                if formato == "pdf":
                    pdf = guardar_en / f"{titulo}.pdf"
                    docx_a_pdf(final, pdf)
                    final = pdf
            if destino == "nextcloud":
                avisos.append(subir_nextcloud(final, carpeta))
            elif destino == "carpeta":
                avisos.append(copiar_a_carpeta(final, carpeta))
            else:
                avisos.append(f"📎 Dejé «{final.name}» en {final.parent}")
            if log:
                log.info("Entrega lista: %s (%s)", final.name, destino)
        except Exception as e:  # noqa: BLE001
            avisos.append(f"⚠️ No pude entregar {pedido_json.stem}: {e}")
            if log:
                log.exception("Falló la entrega %s", pedido_json.name)
            guardar_en = CARPETA / "con_error"
            guardar_en.mkdir(exist_ok=True)
        for f in (pedido_json, CARPETA / f"{pedido_json.stem}.md"):
            if f.exists():
                shutil.move(str(f), str(guardar_en / f.name))
    return avisos
